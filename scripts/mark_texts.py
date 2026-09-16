#!/usr/bin/env python3
"""
Record in texts/ui-szovegek-javitott.xlsx what happened to each suggested
text: a new column ALKALMAZÁS says whether it went into the app unchanged,
was reworded (and why), or was left out (and why). Reviewer notes in
MEGJEGYZÉS get their answer in a VÁLASZ column.
"""
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

P = 'texts/ui-szovegek-javitott.xlsx'
wb = load_workbook(P)
ws = wb.worksheets[0]
hdr = [c.value for c in ws[1]]
if 'ALKALMAZÁS' not in hdr:
    ws.cell(row=1, column=len(hdr) + 1, value='ALKALMAZÁS')
    ws.cell(row=1, column=len(hdr) + 2, value='VÁLASZ')
    hdr += ['ALKALMAZÁS', 'VÁLASZ']
CA = hdr.index('ALKALMAZÁS') + 1
CV = hdr.index('VÁLASZ') + 1
for c in (CA, CV):
    ws.cell(row=1, column=c).font = Font(bold=True)
    ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 48

def norm(s): return (str(s) if s is not None else '').replace('\r\n', '\n').strip()

REWORDED = {
    # id prefix -> reason
    '07.PRESSURE-B.depthsimon.instruction': 'A mechanikát javítottuk: a két gyűrű most CIÁNKÉK (közeli) és MAGENTA (távoli), a gömb színe mondja, melyikhez; a szöveg ehhez igazítva, a lektor stílusában.',
    '07.PRESSURE-B.depthsimon.hint': 'A gyűrűk színe és a leképezés változott (ciánkék = vissza, magenta = előre); a tipp ehhez igazítva.',
}
ANSWERS = {
    '03.NAV.map.instruction': 'Ellenőrizve: a térkép-rész felváltva hely- és iránypróba (4+4), az iránytárcsa mindhárom platformon megjelenik. A javasolt szöveg változatlanul bekerült.',
    '04.REACT-B.': 'A B változat VR-only (a katalógus így hirdeti); a PLATFORM „mind” csak azt jelzi, hogy a szöveg nem ágazik el platformonként. Rendben, nem kell módosítás.',
    '07.PRESSURE-B.': 'A B változat VR-only; a PLATFORM „mind” csak azt jelzi, hogy a szöveg nem ágazik el platformonként. Rendben.',
    '12.FIELD.threshold.instruction.desktop': 'A megvalósítás: KOCKA → SZÓKÖZ vagy bal kattintás, GÖMB → jobb kattintás. A táblázat szövege a helyes, a részletes specifikáció elavult (frissítjük).',
    '12.FIELD.dva.instruction.vr': 'A megvalósítás közvetlen kontrollergombokat használ (nem tárcsát) VR-ban; a táblázat szövege a helyes.',
    '12.FIELD.dva.hint.vr': 'A megvalósítás közvetlen kontrollergombokat használ; a táblázat szövege a helyes.',
    '17.RISK.cards.instruction.mobile': 'A mobil gombok felirata 1 · 2 · 3 · 4 (a pozíció szerint, nem A/B/C/D, hogy ismételt felvételnél ne legyen átvihető a válasz). A táblázat szövege a helyes.',
    '17.RISK.cards.hint.mobile': 'A mobil gombok felirata 1 · 2 · 3 · 4; a táblázat szövege a helyes.',
    '17.RISK.setCardControls.7': 'Egységesítve: a pakligombok neve 1 · 2 · 3 · 4.',
    '16.HANDS.calibrate.1': 'A megvalósítás: a domináns kéz felemelése + ravasz; a táblázat szövege a helyes, a specifikáció elavult.',
    'UI.hub.dom.16': 'Rendben: a HANDS csak VR-ban fut, a javított mondat ezt pontosítja.',
    'UI.hub.dom.17': 'Tudjuk, dinamikus töredék; a javítás bekerült.',
    'UI.hub.dom.18': 'Dinamikus töredék; a javítás bekerült.',
    'UI.domain.20': 'Egységesítve PERFORMANCE INDEX-re.',
}

def svg(old):
    return old.startswith('M') and len(old) > 40 and re.fullmatch(r'[MLCZmlczHhVvSsQqTtAa0-9 .,\-]+', old) is not None

n_same = n_rew = n_skip = 0
for row in ws.iter_rows(min_row=2):
    rid = norm(row[1].value)
    old, new, note = norm(row[8].value), norm(row[9].value), norm(row[10].value)
    a = v = ''
    if svg(old):
        a = 'kihagyva: SVG útvonaladat, nem felhasználói szöveg (az exportból törölve)'
        n_skip += 1
    elif rid.startswith('19.INTENT') and '→' in new:
        a = 'módosítva: a „→” jelek helyett kettőspont/szó, mert ebben a modulban a nyíl maga az irányinger (tesztelői kérés); a szöveg egyébként változatlan'
        n_rew += 1
    elif any(rid.startswith(k) for k in REWORDED):
        a = 'módosítva: ' + next(r for k, r in REWORDED.items() if rid.startswith(k))
        n_rew += 1
    elif old != new:
        a = 'változatlanul beépítve'
        n_same += 1
    if note:
        for k, ans in ANSWERS.items():
            if rid.startswith(k):
                v = ans
                break
    if a: ws.cell(row=row[0].row, column=CA, value=a)
    if v: ws.cell(row=row[0].row, column=CV, value=v)
    for c in (CA, CV):
        ws.cell(row=row[0].row, column=c).alignment = Alignment(wrap_text=True, vertical='top')

wb.save(P)
print(f'változatlanul: {n_same} | módosítva: {n_rew} | kihagyva: {n_skip}')
