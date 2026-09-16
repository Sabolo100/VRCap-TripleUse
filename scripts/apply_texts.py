#!/usr/bin/env python3
"""
Copy corrected texts from texts/ui-szovegek-javitott.xlsx back into the source.

For every row whose JAVÍTOTT SZÖVEG differs from JELENLEGI SZÖVEG, the old text
is located in the file named in FÁJL (or under it, if a directory) as a JS
string literal - including literals split over several lines with `' +` - and
replaced by the new text, re-wrapped in the same style. Rows that cannot be
matched are listed so they can be done by hand.

  python3 scripts/apply_texts.py            # dry run: report only
  python3 scripts/apply_texts.py --write    # apply
"""
import re, sys, os, glob
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'texts/ui-szovegek-javitott.xlsx')
WRITE = '--write' in sys.argv

def norm(s): return (str(s) if s is not None else '').replace('\r\n', '\n').strip()

def js_unescape(body):
    return body.replace("\\'", "'").replace('\\"', '"').replace('\\n', '\n').replace('\\\\', '\\')

def js_escape(text, quote):
    t = text.replace('\\', '\\\\').replace('\n', '\\n')
    return t.replace(quote, '\\' + quote)

# A literal chain: one or more '...'/"..." pieces joined by `+` and whitespace.
PIECE = r"(?:'(?:[^'\\\n]|\\.)*'|\"(?:[^\"\\\n]|\\.)*\")"
CHAIN = re.compile(PIECE + r"(?:\s*\+\s*(?://[^\n]*\n\s*)?" + PIECE + r")*")

def chain_text(chain):
    parts = re.findall(PIECE, chain)
    return ''.join(js_unescape(p[1:-1]) for p in parts)

def rewrap(text, quote, indent, single_ok_width=110):
    """Re-emit `text` as a literal chain in the file's style."""
    if '\n' not in text and len(text) <= single_ok_width:
        return quote + js_escape(text, quote) + quote
    # split into ~95-char pieces at spaces, preserving the pieces' spaces
    words = text.split(' ')
    pieces, cur = [], ''
    for w in words:
        cand = (cur + ' ' + w) if cur else w
        if len(cand) > 95 and cur:
            pieces.append(cur + ' ')
            cur = w
        else:
            cur = cand
    pieces.append(cur)
    joiner = ' +\n' + indent
    return joiner.join(quote + js_escape(p, quote) + quote for p in pieces)

def candidate_files(fajl):
    fajl = (fajl or '').split(':')[0].strip()
    if not fajl:
        return []
    path = os.path.join(ROOT, fajl)
    if os.path.isdir(path):
        return sorted(glob.glob(os.path.join(path, '**/*.ts'), recursive=True))
    return [path] if os.path.exists(path) else []

def find_and_replace(src, old, new):
    """Return (new_src, count) replacing literal chains equal to `old`."""
    out, count, pos = [], 0, 0
    for m in CHAIN.finditer(src):
        if chain_text(m.group(0)) == old:
            quote = m.group(0)[0]
            # indentation of the line the chain starts on
            line_start = src.rfind('\n', 0, m.start()) + 1
            indent = re.match(r'\s*', src[line_start:m.start()]).group(0)
            indent = ' ' * (len(indent) + 2) if '\n' not in m.group(0) else re.search(r'\n(\s*)', m.group(0)).group(1)
            out.append(src[pos:m.start()])
            out.append(rewrap(new, quote, indent))
            pos = m.end()
            count += 1
    out.append(src[pos:])
    return ''.join(out), count

wb = load_workbook(XLSX)
ws = wb.worksheets[0]
rows = list(ws.iter_rows(values_only=True))
header, rows = rows[0], rows[1:]
changed = [r for r in rows if norm(r[8]) != norm(r[9])]

edits = {}      # file -> src
applied, unmatched, multi = [], [], []
for r in changed:
    sor, rid, kat, modul, valt, hely, mezo, plat, old, new, megj, fajl = r[:12]
    old, new = norm(old), norm(new)
    if old.startswith('M') and len(old) > 40 and re.fullmatch(r'[MLCZmlczHhVvSsQqTtAa0-9 .,\-]+', old):
        continue  # SVG path data, not text
    files = candidate_files(fajl)
    hit = None
    for f in files:
        src = edits.get(f) or open(f, encoding='utf-8').read()
        new_src, n = find_and_replace(src, old, new)
        if n:
            hit = (f, n, new_src)
            break
    if not hit:
        unmatched.append((rid, fajl, old[:80]))
        continue
    f, n, new_src = hit
    if n > 1:
        multi.append((rid, f, n))
    edits[f] = new_src
    applied.append((rid, os.path.relpath(f, ROOT)))

print(f'változott sor: {len(changed)} | illesztve: {len(applied)} | nem található: {len(unmatched)} | többszörös: {len(multi)}')
for rid, fajl, old in unmatched:
    print('  ?', rid, '|', fajl, '|', old)
for rid, f, n in multi:
    print('  x'+str(n), rid, os.path.relpath(f, ROOT))
if WRITE:
    for f, src in edits.items():
        open(f, 'w', encoding='utf-8').write(src)
    print('írva:', len(edits), 'fájl')
